"""Этап 3. Данные и библиотеки: слои, импорт с проверками, предпросмотр, транзакция.

Слои (спецификация 10.1–10.4): системный (только чтение) < предприятия (администратор) < пользовательский.
Эффективная запись — из самого приоритетного слоя, где есть ключ.
Импорт: проверки до записи, режимы «добавить / обновить по ключу / заменить слой», предпросмотр,
одна транзакция с откатом. Ошибки построчно; выбор «только корректные строки» или «ничего».
"""
from __future__ import annotations

import csv
import copy
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional, Tuple

LAYER_PRIORITY = {"системный": 0, "предприятия": 1, "пользовательский": 2}


@dataclass
class FieldSpec:
    name: str
    kind: str = "str"          # str | float
    required: bool = False
    min: Optional[float] = None
    gt: Optional[float] = None


@dataclass
class Schema:
    name: str
    key: Tuple[str, ...]
    fields: List[FieldSpec]
    references: Dict[str, str] = field(default_factory=dict)   # поле -> имя библиотеки-ссылки


@dataclass
class Problem:
    row: int
    message: str


@dataclass
class Preview:
    added: List[dict]
    changed: List[dict]
    skipped: List[dict]
    errors: List[Problem]
    keep: List[dict] = field(default_factory=list)   # «заменить»: корректные записи без изменений остаются в слое


class Store:
    def __init__(self):
        self.layers: Dict[str, Dict[str, Dict[Tuple, dict]]] = {}   # layer -> library -> key -> record
        self._snapshot: Optional[dict] = None
        self.mappings: Dict[str, Dict[str, str]] = {}               # схема -> сопоставление колонок (10.4)

    def remember_mapping(self, schema: str, mapping: Dict[str, str]) -> None:
        """Запоминание сопоставления колонок файла с полями схемы (10.4)."""
        self.mappings[schema] = dict(mapping)

    def saved_mapping(self, schema: str) -> Dict[str, str]:
        return dict(self.mappings.get(schema, {}))

    # --- чтение ---
    def get(self, library: str, key: Tuple) -> Optional[dict]:
        best = None
        for layer, libs in self.layers.items():
            rec = libs.get(library, {}).get(key)
            if rec is not None and (best is None or LAYER_PRIORITY[layer] > LAYER_PRIORITY[best[0]]):
                best = (layer, rec)
        return None if best is None else best[1]

    def layer(self, layer: str, library: str) -> Dict[Tuple, dict]:
        return self.layers.setdefault(layer, {}).setdefault(library, {})

    # --- транзакция ---
    def begin(self):
        self._snapshot = copy.deepcopy(self.layers)

    def commit(self):
        self._snapshot = None

    def rollback(self):
        if self._snapshot is not None:
            self.layers = self._snapshot
            self._snapshot = None


def _convert(value: str, f: FieldSpec):
    if value is None or value == "":
        return None
    if f.kind == "float":
        return float(str(value).replace(",", "."))
    return value.strip()


def validate_row(row: dict, schema: Schema, index: int) -> List[Problem]:
    errs: List[Problem] = []
    for f in schema.fields:
        try:
            v = _convert(row.get(f.name, ""), f)
        except ValueError:
            errs.append(Problem(index, f"{f.name}: не число"))
            continue
        if v is None:
            if f.required:
                errs.append(Problem(index, f"{f.name}: обязательное поле пусто"))
            continue
        if f.min is not None and v < f.min:
            errs.append(Problem(index, f"{f.name}: меньше {f.min}"))
        if f.gt is not None and v <= f.gt:
            errs.append(Problem(index, f"{f.name}: должно быть больше {f.gt}"))
    return errs


def parse(rows: List[dict], schema: Schema) -> Tuple[List[dict], List[Problem]]:
    """Преобразование строк в записи + проверки внутри файла (дубликаты ключа)."""
    recs: List[dict] = []
    errs: List[Problem] = []
    seen: Dict[Tuple, int] = {}
    for i, row in enumerate(rows, start=2):   # строка 1 — заголовок
        row_errs = validate_row(row, schema, i)
        rec = {}
        for f in schema.fields:
            try:
                rec[f.name] = _convert(row.get(f.name, ""), f)
            except ValueError:
                rec[f.name] = None
        key = tuple(rec.get(k) for k in schema.key)
        if any(k is None for k in key):
            row_errs.append(Problem(i, "ключ не заполнен"))
        elif key in seen:
            row_errs.append(Problem(i, f"дубликат ключа {key}, уже в строке {seen[key]}"))
        else:
            seen[key] = i
        errs.extend(row_errs)
        recs.append({"_row": i, "_key": key, "_ok": not row_errs, **rec})
    return recs, errs


def check_references(recs: List[dict], schema: Schema, store: Store) -> List[Problem]:
    """Проверка ссылок (10.4). references: поле → библиотека (ключ — значение поля) или
    поле → (библиотека, (поля составного ключа)) — например, оснастка → строка ∅ профиля."""
    errs: List[Problem] = []
    for r in recs:
        for fld, spec in schema.references.items():
            if isinstance(spec, str):
                lib, key_fields = spec, (fld,)
            else:
                lib, key_fields = spec
            values = [r.get(f) for f in key_fields]
            if any(v is None for v in values):
                # необязательная ссылка пуста — не проверяется; обязательная проверяется полем схемы
                continue
            key = tuple(values)
            if store.get(lib, key) is None:
                shown = key[0] if len(key) == 1 else key
                errs.append(Problem(r["_row"], f"{fld} = {shown}: нет записи в библиотеке «{lib}»"))
    return errs


def preview(rows: List[dict], schema: Schema, store: Store, layer: str, mode: str = "обновить") -> Preview:
    """Предпросмотр: «будет добавлено, изменено, пропущено». mode: добавить | обновить | заменить."""
    recs, errs = parse(rows, schema)
    errs += check_references(recs, schema, store)
    bad_rows = {p.row for p in errs}
    added, changed, skipped, keep = [], [], [], []
    target = store.layers.get(layer, {}).get(schema.name, {})     # без побочного создания слоя
    if mode == "заменить":
        existing_keys = set(target.keys())
        new_keys = {r["_key"] for r in recs if r["_ok"] and r["_row"] not in bad_rows}
        for k in existing_keys - new_keys:
            skipped.append({"key": k, "reason": "удалится при замене слоя"})
    for r in recs:
        if r["_row"] in bad_rows:
            skipped.append({"row": r["_row"], "reason": "ошибка"})
            continue
        current = target.get(r["_key"])
        data = {k: v for k, v in r.items() if not k.startswith("_")}
        if current is None:
            added.append(data)
        elif mode == "добавить":
            skipped.append({"row": r["_row"], "reason": "ключ уже есть"})
        elif current != data:
            changed.append(data)
        else:
            skipped.append({"row": r["_row"], "reason": "без изменений"})
            keep.append(data)
    return Preview(added, changed, skipped, errs, keep)


def apply(rows: List[dict], schema: Schema, store: Store, layer: str,
          mode: str = "обновить", only_ok: bool = True) -> Preview:
    """Применение в одной транзакции. only_ok=False — если есть ошибки, ничего не пишем."""
    pv = preview(rows, schema, store, layer, mode)
    if pv.errors and not only_ok:
        return pv
    store.begin()
    try:
        target = store.layer(layer, schema.name)
        if mode == "заменить":
            target.clear()
        for rec in pv.added + pv.changed + pv.keep:
            key = tuple(rec[k] for k in schema.key)
            target[key] = rec
        store.commit()
    except Exception:
        store.rollback()
        raise
    return pv


# ---------- схемы для проверки цепочек ссылок (10.4) ----------
# Оснастка (2.2.2): позиция × профиль × ∅. Ссылка «позиция» — в библиотеке позиций оборудования;
# «профиль + ∅» — строка ∅ профиля трубогиба (библиотека «строки профилей», ключ (профиль, ∅)).
TOOL_SCHEMA = Schema(
    name="оснастка",
    key=("позиция", "профиль", "od_mm"),
    fields=[
        FieldSpec("позиция", "str", required=True),
        FieldSpec("профиль", "str", required=True),
        FieldSpec("od_mm", "float", required=True, gt=0),
        FieldSpec("инв_номер", "str"),
    ],
    references={"позиция": "позиции оборудования",
                "профиль": ("строки профилей", ("профиль", "od_mm"))},
)

# Пресет (2.3): назначение (номенклатура), трубогиб, инструмент навивки (необязательно), изоляция и система (необязательно).
PRESET_SCHEMA = Schema(
    name="пресеты",
    key=("имя",),
    fields=[
        FieldSpec("имя", "str", required=True),
        FieldSpec("назначение", "str", required=True),
        FieldSpec("трубогиб", "str", required=True),
        FieldSpec("инструмент навивки", "str"),
        FieldSpec("режим заготовки", "str"),
        FieldSpec("компенсация", "str"),
        FieldSpec("изоляция", "str"),
        FieldSpec("система", "str"),
    ],
    references={"назначение": "номенклатура",
                "трубогиб": "трубогибы",
                "инструмент навивки": "трубогибы",
                "изоляция": "изоляция",
                "система": "системы"},
)


def load_csv(path: str) -> List[dict]:
    with open(path, encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def load_xlsx(path: str) -> List[dict]:
    """Первый лист XLSX: первая строка — заголовки; значения — строками (как из CSV)."""
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if header is None:
        return []
    names = [str(h).strip() if h is not None else "" for h in header]
    out: List[dict] = []
    for r in rows_iter:
        if r is None or all(v is None or v == "" for v in r):
            continue
        out.append({n: ("" if v is None else str(v)) for n, v in zip(names, r) if n})
    wb.close()
    return out


# единицы колонок при импорте (10.4): мм — как есть, дюймы → мм, кг/м³ — без пересчёта
UNIT_FACTORS = {"мм": 1.0, "дюймы": 25.4, "кг/м³": 1.0}


def convert_units(rows: List[dict], units: Dict[str, str]) -> List[dict]:
    """Пересчёт колонок с единицами (дюймы → мм). Неизвестная единица — ValueError."""
    out: List[dict] = []
    for row in rows:
        new = dict(row)
        for col, unit in units.items():
            if unit not in UNIT_FACTORS:
                raise ValueError(f"неизвестная единица колонки {col}: {unit}")
            v = row.get(col, "")
            if v not in (None, ""):
                new[col] = str(float(str(v).replace(",", ".")) * UNIT_FACTORS[unit])
        out.append(new)
    return out


def map_columns(rows: List[dict], mapping: Dict[str, str]) -> List[dict]:
    """Переименование колонок файла в поля схемы; колонки без сопоставления отбрасываются (10.4)."""
    return [{mapping[k]: v for k, v in row.items() if k in mapping} for row in rows]


DIAMETER_SCHEMA = Schema(
    name="диаметры",
    key=("od_mm", "group"),
    fields=[
        FieldSpec("od_mm", "float", required=True, gt=0),
        FieldSpec("group", "str", required=True),
        FieldSpec("R_mm", "float", required=True, gt=0),
        FieldSpec("Lmin_mm", "float", required=True, min=0),
        FieldSpec("N_mm", "float", required=True, min=0),
        FieldSpec("Ls_mm", "float", min=0),
        FieldSpec("inch", "str"),
        FieldSpec("flag", "str"),
    ],
)
